from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from audit_data import FINDINGS, POSITIVE_CONTROLS


OUTPUT = Path(__file__).with_name("Laporan_Audit_Keamanan_SIMANTAP_2026-10-18.docx")
NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F4F6F8"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
SLATE = RGBColor(71, 85, 105)


def font(run, size=10.5, bold=False, color=BLACK, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    el = tc_pr.find(qn("w:shd"))
    if el is None:
        el = OxmlElement("w:shd")
        tc_pr.append(el)
    el.set(qn("w:fill"), fill)


def margins(cell, top=105, start=120, bottom=105, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def borders(table):
    tbl_pr = table._tbl.tblPr
    group = tbl_pr.first_child_found_in("w:tblBorders")
    if group is None:
        group = OxmlElement("w:tblBorders")
        tbl_pr.append(group)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = group.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            group.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def cell_text(cell, value, bold=False, white=False, center=False, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(value))
    font(run, size=size, bold=bold, color=RGBColor(255, 255, 255) if white else BLACK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)


def heading(doc, text, level=1, page=False):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page
    return p


def body(doc, text, lead=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if lead:
        r = p.add_run(lead)
        font(r, bold=True)
    r = p.add_run(text)
    font(r)
    r.italic = italic
    return p


def bullets(doc, items, numbered=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    p.add_run().add_break()
    for idx, item in enumerate(items, start=1):
        prefix = f"{idx}.  " if numbered else "•  "
        run = p.add_run(f"{prefix}{item}")
        font(run, size=10.1)
        if idx < len(items):
            run.add_break()


def label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(text), bold=True)


def labeled_bullets(doc, title, items):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    font(p.add_run(title), bold=True)
    for item in items:
        p.add_run().add_break()
        font(p.add_run(f"•  {item}"), size=10.1)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    font(run, size=8.2, color=RGBColor(100, 116, 139))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def add_finding(doc, item):
    heading(doc, f"{item['id']}  {item['title']}", 2, page=item["id"] == "F13")
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(5.25)
    cell_text(table.cell(0, 0), "Tingkat risiko", bold=True, white=True)
    cell_text(table.cell(0, 1), f"{item['severity']}  {item['priority']}", bold=True, white=True)
    shade(table.cell(0, 0), NAVY)
    shade(table.cell(0, 1), NAVY)
    repeat_header(table.rows[0])
    cell_text(table.cell(1, 0), "Klasifikasi", bold=True)
    cell_text(table.cell(1, 1), item["cwe"])
    borders(table)
    for row in table.rows:
        no_split(row)

    body(doc, item["why"], "Mengapa tidak aman. ")
    body(doc, item["impact"], "Skenario dampak. ")
    labeled_bullets(doc, "Bukti kode", item["evidence"])
    labeled_bullets(doc, "Cara penanganan", item["fix"])
    labeled_bullets(doc, "Kriteria verifikasi", item["verify"])


def setup(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 18, 14, 7),
        ("Heading 2", 13.5, 11, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.size = Pt(10.1)

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(hp.add_run("AUDIT KEAMANAN SIMANTAP"), size=8, bold=True, color=RGBColor(100, 116, 139))
    add_page_field(section.footer.paragraphs[0])


def cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(68)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(12)
    p.add_run("Laporan Audit Keamanan Aplikasi SIMANTAP")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    font(
        p.add_run("Daftar temuan celah keamanan alasan risiko dan rencana penanganan"),
        size=14,
        color=RGBColor(51, 65, 85),
    )
    label(doc, "Kesimpulan audit")
    body(
        doc,
        "Risiko keseluruhan dinilai KRITIS. Aplikasi memiliki kontrol dasar yang baik, tetapi jalur autentikasi lokal dapat melewati validasi eGov atau SIMPEG, kredensial sensitif tersimpan dalam kode dan paket deploy, serta dependency memiliki advisory critical dan high. Sistem belum layak dianggap aman untuk produksi sebelum tindakan P0 selesai dan diverifikasi.",
    )
    for key, value in (
        ("Objek audit", "Project SIMANTAP di /Users/simplephi/Documents/riswan/simantap"),
        ("Tanggal pemeriksaan", "18 Oktober 2026"),
        ("Metode", "Static application security review dan software composition analysis"),
        ("Klasifikasi", "Internal dan rahasia"),
        ("Status", "Laporan awal berbasis kode dan artefak yang tersedia"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        font(p.add_run(f"{key}: "), bold=True)
        font(p.add_run(value))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("Catatan kerahasiaan: nilai kata sandi, connection string, dan secret tidak disalin ke laporan.")
    font(r, size=9.3, color=SLATE)
    r.italic = True
    doc.add_page_break()


def summary(doc):
    heading(doc, "Ringkasan Eksekutif", 1)
    body(
        doc,
        "Audit menemukan 13 isu: 4 critical, 4 high, 3 medium, dan 2 low. Tiga temuan pertama dapat menyebabkan pengambilalihan akun atau pemalsuan sesi tanpa harus merusak eGov atau SIMPEG. Kelemahan tersebut berada dalam kode SIMANTAP dan tetap menjadi tanggung jawab aplikasi ini.",
    )
    body(
        doc,
        "Ketergantungan pada akun eGov dan SIMPEG diterima sebagai batas arsitektur. Audit tidak menilai kebijakan kata sandi, integritas data, atau ketersediaan sistem eksternal. Audit tetap menilai cara SIMANTAP memanggil sistem tersebut, menyimpan hasil autentikasi, melindungi secret, dan membatasi percobaan login.",
    )
    labeled_bullets(doc, "Kontrol yang sudah ada dan perlu dipertahankan:", POSITIVE_CONTROLS)
    heading(doc, "Prioritas Penanganan", 2)
    bullets(
        doc,
        [
            "Dalam 24 jam: nonaktifkan autentikasi lokal untuk akun eGov atau SIMPEG, hapus password default, rotasi kredensial database dan JWT, serta tarik paket deploy berisi environment.",
            "Dalam 72 jam: perbarui dependency critical dan high melalui branch terpisah, batasi Swagger dan file upload, lalu aktifkan rate limiting login dan refresh.",
            "Dalam 7 hari: pindahkan refresh token ke cookie HttpOnly, terapkan security headers dan HTTPS penuh, serta tambah regression test autentikasi dan otorisasi.",
            "Dalam 30 hari: lengkapi monitoring, uji penetrasi staging, dan buktikan akun MySQL benar-benar SELECT-only pada tingkat database.",
        ],
        numbered=True,
    )

    heading(doc, "Ringkasan Temuan", 2, page=True)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.55, 0.85, 4.65, 0.85]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for i, value in enumerate(("ID", "Risiko", "Temuan", "Prioritas")):
        cell_text(table.cell(0, i), value, bold=True, white=True, center=True)
        shade(table.cell(0, i), NAVY)
    repeat_header(table.rows[0])
    for idx, item in enumerate(FINDINGS):
        cells = table.add_row().cells
        vals = (item["id"], item["severity"].title(), item["title"], item["priority"])
        for i, value in enumerate(vals):
            cell_text(cells[i], value, bold=i in (0, 1, 3), center=i != 2)
            if idx % 2:
                shade(cells[i], PALE_BLUE)
        no_split(table.rows[-1])
    borders(table)


def scope(doc):
    heading(doc, "Ruang Lingkup dan Batasan", 1)
    body(doc, "Pemeriksaan mencakup komponen berikut.")
    bullets(
        doc,
        [
            "Backend NestJS: autentikasi, otorisasi, integrasi eGov atau SIMPEG, upload, logging, CORS, Swagger, dan health check.",
            "Frontend Next.js: penyimpanan token, alur refresh, konfigurasi deploy statis, dan komunikasi API.",
            "Prisma schema, migration, seed, konfigurasi, package manifest, lockfile, dan arsip deploy di direktori project.",
            "Software composition analysis memakai npm audit terhadap dependency produksi pada 18 Oktober 2026.",
        ],
    )
    body(doc, "Pemeriksaan tidak mencakup hal berikut.")
    bullets(
        doc,
        [
            "Penetration test terhadap server produksi, database eGov atau SIMPEG, maupun layanan SiRUP.",
            "Verifikasi reverse proxy, firewall, WAF, sertifikat TLS, permission filesystem produksi, backup, dan monitoring infrastruktur.",
            "Phishing, keamanan perangkat pengguna, atau kebijakan internal pengelola akun eGov dan SIMPEG.",
            "Jaminan privilege MySQL. Kode yang diperiksa hanya menjalankan SELECT, tetapi hak database harus dibuktikan di server.",
        ],
    )
    body(
        doc,
        "Risiko menggabungkan kemudahan eksploitasi, tingkat akses yang diperoleh, cakupan data, dan dampak operasional. Temuan yang bergantung pada konfigurasi produksi diberi batasan karena bukti runtime tidak tersedia.",
    )


def remediation_plan(doc):
    heading(doc, "Rencana Remediasi Bertahap", 1, page=True)
    rows = [
        ("P0", "0-24 jam", "Aplikasi dan infrastruktur", "F01-F05", "Tutup bypass login, rotasi secret, tarik artefak bocor, dan mulai patch"),
        ("P1", "1-7 hari", "Backend frontend DevOps", "F06-F09", "Cookie HttpOnly, upload privat, rate limit, HTTPS, dan headers"),
        ("P2", "1-2 minggu", "Backend dan keamanan", "F10-F11", "Batasi dokumentasi, logging terstruktur, alert, dan retensi"),
        ("P3", "Maks 30 hari", "Backend dan product owner", "F12-F13", "Kurangi informasi publik dan perketat schema input"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, w in enumerate((0.65, 0.9, 1.5, 0.75, 3.1)):
        table.columns[i].width = Inches(w)
    for i, value in enumerate(("Tahap", "Waktu", "Penanggung jawab", "Temuan", "Hasil")):
        cell_text(table.cell(0, i), value, bold=True, white=True, center=True)
        shade(table.cell(0, i), NAVY)
    repeat_header(table.rows[0])
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, bold=i == 0, center=i in (0, 1, 3), size=8.8)
            if idx % 2:
                shade(cells[i], PALE_BLUE)
        no_split(table.rows[-1])
    borders(table)

    heading(doc, "Desain Autentikasi yang Direkomendasikan", 2)
    body(
        doc,
        "Akun pegawai tetap berasal dari eGov atau SIMPEG. SIMANTAP hanya menyimpan profil, status akses, role, scope OPD, dan session token miliknya sendiri. Keputusan benar atau salahnya kata sandi selalu berasal dari eGov atau SIMPEG untuk authSource EGOV.",
    )
    bullets(
        doc,
        [
            "Pengguna mengirim identifier dan kata sandi melalui HTTPS ke endpoint SIMANTAP.",
            "SIMANTAP menerapkan rate limit lalu memverifikasi kredensial ke eGov atau SIMPEG melalui koneksi read-only atau layanan resmi.",
            "Jika valid, SIMANTAP mengecek akun lokal untuk status AKTIF, role, dan scope. Akun tanpa role ditolak tanpa membuat kredensial lokal.",
            "SIMANTAP menerbitkan access token singkat dan refresh token acak melalui cookie HttpOnly. Hash kata sandi eGov tidak disimpan.",
            "Perubahan role atau pencabutan akses mencabut seluruh refresh token. Validasi user aktif tetap dilakukan server-side.",
        ],
        numbered=True,
    )

    heading(doc, "Pengujian Wajib Setelah Perbaikan", 2)
    rows = [
        ("Autentikasi", "Password default dan password eGov lama ditolak; akun nonaktif tidak mendapat token"),
        ("Otorisasi", "Setiap role diuji pada endpoint read dan write serta scope OPD dan subunit"),
        ("Session", "Refresh rotation, logout, pencabutan role, reuse token, dan expiry berjalan benar"),
        ("Upload", "Spoofed extension, polyglot, oversize, unauthorized download, dan malware test ditolak"),
        ("Transport", "HTTP redirect, CORS production, CSP, HSTS, nosniff, dan frame protection aktif"),
        ("Dependency", "SCA tidak memiliki critical dan setiap high memiliki perbaikan atau penerimaan risiko"),
        ("Logging", "Password spraying dan perubahan role menghasilkan alert tanpa secret atau data pribadi berlebih"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.65)
    for i, value in enumerate(("Area", "Kriteria lulus")):
        cell_text(table.cell(0, i), value, bold=True, white=True, center=True)
        shade(table.cell(0, i), NAVY)
    repeat_header(table.rows[0])
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        cell_text(cells[0], row[0], bold=True)
        cell_text(cells[1], row[1])
        if idx % 2:
            shade(cells[0], PALE_BLUE)
            shade(cells[1], PALE_BLUE)
        no_split(table.rows[-1])
    borders(table)


def conclusion(doc):
    heading(doc, "Kesimpulan Auditor", 1, page=True)
    body(
        doc,
        "SIMANTAP belum dapat dinyatakan aman untuk produksi dalam kondisi yang diperiksa. Risiko utama berasal dari bypass validasi eGov melalui kata sandi lokal, penggunaan kata sandi default, secret dalam kode dan arsip deploy, serta dependency dengan advisory critical. Area tersebut dapat memungkinkan pengambilalihan akun, pemalsuan sesi, akses database, atau gangguan layanan.",
    )
    body(
        doc,
        "Penggunaan akun eGov atau SIMPEG bukan temuan dengan sendirinya. Integrasi dapat dipertahankan setelah SIMANTAP berhenti menyimpan atau menerima kata sandi eksternal secara lokal, memakai akun database SELECT-only, melindungi secret, dan membatasi percobaan.",
    )
    body(
        doc,
        "Status risiko dapat diturunkan setelah seluruh P0 ditutup, bukti rotasi secret tersedia, regression test lulus, dan audit dependency tidak lagi menunjukkan critical. Penetration test pada staging yang menyerupai produksi tetap diperlukan sebelum persetujuan akhir.",
    )

    heading(doc, "Lampiran Bukti Utama", 1)
    rows = [
        ("Autentikasi lokal", "backend/src/modules/auth/auth.service.ts", "75-218"),
        ("Role dan default password", "backend/src/modules/users/users.service.ts", "1102-1273"),
        ("Seed administrator", "backend/prisma/seed.ts", "6-124"),
        ("Kredensial eGov", "backend/src/config/egov.config.ts", "1-10"),
        ("Autentikasi eGov", "backend/src/core/egov/egov.service.ts", "139-310"),
        ("Konfigurasi server", "backend/src/main.ts", "14-103"),
        ("Upload bukti", "backend/src/modules/pembangunan/pembangunan.controller.ts", "359-413"),
        ("Penyimpanan token", "frontend/src/lib/token-storage.ts", "1-70"),
        ("Dependency backend", "backend/package.json dan package-lock.json", "audit 18-10-2026"),
        ("Dependency frontend", "frontend/package.json dan package-lock.json", "audit 18-10-2026"),
        ("Arsip deploy", "monalisa-backend.zip", ".env ditemukan"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, w in enumerate((1.6, 4.4, 1.0)):
        table.columns[i].width = Inches(w)
    for i, value in enumerate(("Area", "Sumber", "Lokasi")):
        cell_text(table.cell(0, i), value, bold=True, white=True, center=True)
        shade(table.cell(0, i), NAVY)
    repeat_header(table.rows[0])
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, bold=i == 0, size=8.6)
            if idx % 2:
                shade(cells[i], PALE_BLUE)
        no_split(table.rows[-1])
    borders(table)

    heading(doc, "Skala Risiko", 2)
    bullets(
        doc,
        [
            "Critical: dapat memberi akses administrator, memalsukan sesi, membuka secret inti, atau memengaruhi banyak pengguna tanpa prasyarat berat.",
            "High: dapat mengambil alih sesi, membuka data lintas kewenangan, mengunggah konten berbahaya, atau menimbulkan gangguan besar.",
            "Medium: kontrol pertahanan tidak lengkap dan memperbesar dampak serangan atau menghambat deteksi.",
            "Low: paparan terbatas yang tetap perlu ditutup sebagai pertahanan berlapis.",
        ],
    )


def build():
    doc = Document()
    setup(doc)
    cover(doc)
    summary(doc)
    scope(doc)
    heading(doc, "Temuan Rinci", 1, page=True)
    for item in FINDINGS:
        add_finding(doc, item)
    remediation_plan(doc)
    conclusion(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
